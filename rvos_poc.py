"""
RVOS Proof of Concept -- core reasoning test.

Two-agent pipeline: extract the paper's claim, then judge its novelty
against related work retrieved from OpenAlex. Orchestrated with LangGraph
as a thin sequencing layer (see build_graph()) -- the agents' reasoning,
retries, and error handling live entirely in the plain functions below.
Deliberately no UI, no database, no batch/multi-paper processing.

Usage:
    python rvos_poc.py path/to/paper.(txt|pdf|docx)

Requires ANTHROPIC_API_KEY in the environment (see .env.example).
"""

import json
import os
import sys
import time
from typing import TypedDict

import pdfplumber
import requests
from anthropic import Anthropic
from docx import Document
from dotenv import load_dotenv
from langgraph.graph import END, START, StateGraph

load_dotenv()

client = Anthropic()  # reads ANTHROPIC_API_KEY from the environment
MODEL = "claude-sonnet-5"  # cheap and sufficient for this step; only escalate if judgment quality is weak in your reading
REQUIRED_KEYS = {"claim", "method", "result", "keywords"}
# Only this many leading characters of a paper reach the model (D1): the report and the UI say so.
MAX_PAPER_CHARS = 12000
# Verdict length asked of the judge (F7). Shipped at the validated 150-250: a
# 60-100 range was tried and changed the judgment on the novel fixture (it said
# "overlaps significantly"). Change only with a live before/after check.
VERDICT_WORDS_MIN = 150
VERDICT_WORDS_MAX = 250

def _response_text(resp) -> str:
    """Sonnet 5 can prepend a ThinkingBlock before the text block, so
    content[0] isn't reliably the answer -- find the first text block."""
    for block in resp.content:
        if block.type == "text":
            return block.text.strip()
    raise ValueError(f"No text block in response: {resp.content!r}")


def extract_claim(paper_text: str) -> dict:
    """Agent 1: pull out the core claim, method, and result.

    Occasionally returns a malformed JSON string (e.g. an invalid escape
    like \\' inside a value), a JSON value that isn't an object, or one
    missing/malformed keywords -- retry a couple of times rather than fail
    the whole pipeline on a single bad sample."""
    prompt = f"""Read this research paper text and extract, in your own words:
1. The core claim (one or two sentences)
2. The method used (one or two sentences)
3. The key stated result (one or two sentences)
4. Three to five search-friendly keyword phrases for finding related work

Return ONLY valid JSON with keys: claim, method, result, keywords (a list of strings).
No markdown fences, no commentary, no escaped quotes inside string values --
just the JSON object.

PAPER TEXT:
{paper_text[:MAX_PAPER_CHARS]}
"""
    last_error = None
    for attempt in range(3):
        resp = client.messages.create(
            model=MODEL,
            max_tokens=600,
            messages=[{"role": "user", "content": prompt}],
        )
        text = _response_text(resp)
        text = text.replace("```json", "").replace("```", "").strip()
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError as e:
            last_error = e
            continue
        if not isinstance(parsed, dict):
            last_error = ValueError(f"Model response was valid JSON but not an object: {parsed!r}")
            continue
        missing = REQUIRED_KEYS - parsed.keys()
        if missing:
            last_error = ValueError(f"Model response missing required keys: {sorted(missing)}")
            continue
        keywords = parsed.get("keywords")
        if not isinstance(keywords, list) or not keywords or not all(
            isinstance(k, str) for k in keywords
        ):
            last_error = ValueError(
                f"Model response's keywords must be a non-empty list of strings, got: {keywords!r}"
            )
            continue
        return parsed
    raise last_error


def _reconstruct_abstract(inverted_index):
    """OpenAlex stores abstracts as an inverted index -- rebuild plain text."""
    if not inverted_index:
        return ""
    positions = {}
    for word, idxs in inverted_index.items():
        for i in idxs:
            positions[i] = word
    return " ".join(positions[i] for i in sorted(positions))


def search_openalex(keywords: list, limit: int = 8) -> list:
    """Evidence retrieval: query OpenAlex for related work. Not an LLM call --
    OpenAlex already does the search; we just ask it."""
    query = " ".join(keywords)
    url = "https://api.openalex.org/works"
    params = {"search": query, "per-page": limit, "sort": "relevance_score:desc"}
    mailto = os.environ.get("OPENALEX_MAILTO")
    if mailto:
        params["mailto"] = mailto
    api_key = os.environ.get("OPENALEX_API_KEY")
    # Sent as an Authorization header, not a query param: a query param ends up
    # in requests' own error message ("... for url: ..."), which leaked the key
    # into a live test's traceback (see the findings register, F17 correction).
    headers = {"Authorization": f"Bearer {api_key}"} if api_key else {}

    # 429/500/503/504 are treated as transient (F16/F17): OpenAlex's own heavy-
    # load 503s were observed with a Retry-After header; 500 is retried the same
    # way by PyAlex, an established OpenAlex client; 504 (Cloudflare timing out
    # waiting on OpenAlex's own backend) was directly observed too. Retry-After
    # is honored when present (capped at 30s -- it sent 60 once; don't trust an
    # arbitrarily large value), otherwise fixed exponential backoff as before.
    for attempt in range(3):
        r = requests.get(url, params=params, headers=headers, timeout=20)
        if r.status_code in (429, 500, 503, 504) and attempt < 2:
            retry_after = r.headers.get("Retry-After")
            try:
                wait = min(float(retry_after), 30) if retry_after is not None else 2 ** attempt
            except ValueError:
                wait = 2 ** attempt  # Retry-After can legally be an HTTP-date; we don't parse that form
            time.sleep(wait)
            continue
        r.raise_for_status()
        break
    results = []
    for w in r.json().get("results", []):
        results.append({
            "title": w.get("title"),
            "year": w.get("publication_year"),
            "id": w.get("id"),
            "abstract": _reconstruct_abstract(w.get("abstract_inverted_index")),
        })
    return results


def judge_novelty(extracted: dict, related_works: list) -> str:
    """Agent 2: the core reasoning step -- compare the claim against retrieved evidence.
    This is the step the whole POC exists to test."""
    evidence_block = "\n\n".join(
        f"[{i+1}] {w['title']} ({w['year']})\n{w['abstract'][:500]}"
        for i, w in enumerate(related_works) if w["abstract"]
    )
    prompt = f"""You are assessing the novelty of a research claim against related published work.

CLAIM: {extracted['claim']}
METHOD: {extracted['method']}
RESULT: {extracted['result']}

RELATED WORK RETRIEVED FROM OPENALEX:
{evidence_block if evidence_block else "(No abstracts were retrievable for the top matches.)"}

Instructions:
- Judge whether the claim appears novel, overlaps significantly with specific retrieved work, or whether there is insufficient evidence to judge.
- If you say something overlaps, name the SPECIFIC numbered source it overlaps with. Never make a vague claim without pointing to a numbered source.
- If the retrieved evidence is thin, unrelated, or abstracts are empty, say "insufficient evidence" explicitly rather than guessing. This is the single most important instruction in this prompt -- do not fabricate a confident verdict from weak evidence.
- Write {VERDICT_WORDS_MIN}-{VERDICT_WORDS_MAX} words of plain prose, not JSON.
"""
    resp = client.messages.create(
        model=MODEL,
        max_tokens=1024,
        messages=[{"role": "user", "content": prompt}],
    )
    return _response_text(resp)


def numbered_related(related: list) -> list:
    """Citation numbers (from 1, in list order) shared by the report and the web response.
    judge_novelty numbers its evidence the same way, so a Verdict's [n] matches both
    (pinned by test_judge_prompt_numbers_match_numbered_related)."""
    return [(i + 1, w) for i, w in enumerate(related)]


class PipelineState(TypedDict):
    """LangGraph state carried through the pipeline. Orchestration only --
    the three functions above keep their own reasoning, retries, and error
    handling unchanged; the graph just sequences them."""
    paper_text: str
    extracted: dict
    related: list
    verdict: str


def _extract_node(state: PipelineState) -> dict:
    print("Extracting claim...")
    return {"extracted": extract_claim(state["paper_text"])}


def _search_node(state: PipelineState) -> dict:
    print("Searching OpenAlex for related work...")
    return {"related": search_openalex(state["extracted"]["keywords"])}


def _judge_node(state: PipelineState) -> dict:
    print("Judging novelty...")
    return {"verdict": judge_novelty(state["extracted"], state["related"])}


def build_graph():
    """Wires extract -> search -> judge as a linear LangGraph pipeline."""
    graph = StateGraph(PipelineState)
    graph.add_node("extract", _extract_node)
    graph.add_node("search", _search_node)
    graph.add_node("judge", _judge_node)
    graph.add_edge(START, "extract")
    graph.add_edge("extract", "search")
    graph.add_edge("search", "judge")
    graph.add_edge("judge", END)
    return graph.compile()


class PaperLoadError(ValueError):
    """The paper file couldn't be turned into usable paper text."""


def _load_txt(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        with open(path, "r", encoding="cp1252") as f:
            return f.read()


def _load_pdf(path: str) -> str:
    try:
        with pdfplumber.open(path) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as e:
        # Deliberately broad: pdfplumber/pdfminer raise their own internal
        # exception types for a corrupt or password-protected PDF, and
        # enumerating those here would be as fragile as the leak this fixes
        # in app.py (see the fix plan, F2). Any failure at this point means
        # the PDF couldn't be read, which is exactly what PaperLoadError is
        # for.
        raise PaperLoadError(
            f"{path} could not be read as a PDF (it may be corrupted or "
            f"password-protected): {e}"
        ) from e


def _load_docx(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


_LOADERS = {".txt": _load_txt, ".pdf": _load_pdf, ".docx": _load_docx}


def load_paper_text(paper_path: str) -> str:
    """Load a .txt, .pdf or .docx paper into plain paper text (no model involved).

    Raises PaperLoadError for an unsupported extension, or when the file has
    no extractable text (e.g. a scanned PDF -- OCR is not supported)."""
    ext = os.path.splitext(paper_path)[1].lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise PaperLoadError(
            f"Unsupported file type {ext!r} for {paper_path}: expected .txt, .pdf or .docx"
        )
    text = loader(paper_path)
    if not text.strip():
        raise PaperLoadError(
            f"{paper_path} has no extractable text (scanned PDFs need OCR, which is not supported)"
        )
    return text


def run(paper_path: str):
    paper_text = load_paper_text(paper_path)

    result = build_graph().invoke({"paper_text": paper_text})
    extracted = result["extracted"]
    related = result["related"]
    verdict = result["verdict"]

    related_lines = "\n".join(f"[{n}] {w['title']} ({w['year']})" for n, w in numbered_related(related))
    truncation_note = (
        f"\n> Note: this paper has {len(paper_text)} characters; only the first {MAX_PAPER_CHARS} were assessed.\n"
        if len(paper_text) > MAX_PAPER_CHARS else ""
    )
    report = f"""# RVOS POC report -- {os.path.basename(paper_path)}
{truncation_note}
## Extracted claim
{extracted['claim']}

## Method
{extracted['method']}

## Stated result
{extracted['result']}

## Related work retrieved ({len(related)} results)
{related_lines}

## Novelty verdict
{verdict}
"""
    out_path = paper_path.rsplit(".", 1)[0] + "_report.md"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(report)
    print(f"\nDone. Report written to {out_path}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Usage: python rvos_poc.py path/to/paper.(txt|pdf|docx)")
        sys.exit(1)
    try:
        run(sys.argv[1])
    except PaperLoadError as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)
